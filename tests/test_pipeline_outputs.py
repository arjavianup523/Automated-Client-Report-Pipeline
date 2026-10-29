
"""End-to-end tests for the automated client reporting pipeline."""

import sqlite3
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "reporting.db"
OUTPUT_DIR = ROOT / "outputs"


def latest_file(extension):
    files = list(OUTPUT_DIR.glob(f"client_report_*.{extension}"))
    assert files, f"No .{extension} report found."
    return max(files, key=lambda path: path.stat().st_mtime)


def get_database_data():
    with sqlite3.connect(DB_PATH) as con:
        con.row_factory = sqlite3.Row

        summary = con.execute("""
            SELECT c.client_name AS client,
                   f.fund_name AS fund,
                   f.benchmark_name AS benchmark,
                   rp.period_start, rp.period_end,
                   COUNT(h.holding_id) AS number_of_holdings,
                   SUM(h.market_value) AS total_market_value,
                   SUM(h.weight) AS total_weight,
                   p.fund_return, p.benchmark_return,
                   p.fund_return - p.benchmark_return AS active_return
            FROM clients c
            JOIN funds f ON f.fund_id = 1
            JOIN reporting_periods rp ON rp.period_id = 1
            JOIN holdings h ON h.fund_id = f.fund_id
                           AND h.period_id = rp.period_id
            JOIN performance p ON p.fund_id = f.fund_id
                              AND p.period_id = rp.period_id
            WHERE c.client_id = 1
            GROUP BY c.client_name, f.fund_name, f.benchmark_name,
                     rp.period_start, rp.period_end,
                     p.fund_return, p.benchmark_return
        """).fetchone()

        holdings = con.execute("""
            SELECT security_id, security_name, sector, market_value, weight
            FROM holdings
            WHERE fund_id = 1 AND period_id = 1
            ORDER BY market_value DESC
        """).fetchall()

    assert summary is not None
    assert holdings
    return dict(summary), [dict(row) for row in holdings]


def extract_presentation_text(presentation):
    content = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                content.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    content.extend(cell.text for cell in row.cells)
    return "\n".join(content)


def test_excel_report_has_expected_worksheets():
    workbook = load_workbook(latest_file("xlsx"), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Executive Summary", "Holdings", "Validation Checks"
        ]
        assert workbook["Executive Summary"].max_row == 12
        assert workbook["Holdings"].max_row >= 2
        assert workbook["Validation Checks"].max_row == 4
    finally:
        workbook.close()


def test_excel_summary_matches_database():
    summary, _ = get_database_data()
    workbook = load_workbook(latest_file("xlsx"), read_only=True, data_only=True)
    try:
        sheet = workbook["Executive Summary"]
        actual = {
            sheet.cell(row, 1).value: sheet.cell(row, 2).value
            for row in range(2, sheet.max_row + 1)
        }
        expected = {
            "Client": summary["client"],
            "Fund": summary["fund"],
            "Benchmark": summary["benchmark"],
            "Period Start": summary["period_start"],
            "Period End": summary["period_end"],
            "Number of Holdings": summary["number_of_holdings"],
            "Total Market Value (INR)": summary["total_market_value"],
            "Total Portfolio Weight": summary["total_weight"],
            "Fund Return": summary["fund_return"],
            "Benchmark Return": summary["benchmark_return"],
            "Active Return": summary["active_return"],
        }
        for key, value in expected.items():
            assert actual[key] == value, f"Incorrect Excel value: {key}"
    finally:
        workbook.close()


def test_excel_holdings_match_database():
    _, expected = get_database_data()
    workbook = load_workbook(latest_file("xlsx"), read_only=True, data_only=True)
    try:
        sheet = workbook["Holdings"]
        headers = [cell.value for cell in sheet[1]]
        actual = [
            dict(zip(headers, row))
            for row in sheet.iter_rows(min_row=2, values_only=True)
        ]
        assert len(actual) == len(expected)

        for row, source in zip(actual, expected):
            assert row["Security ID"] == source["security_id"]
            assert row["Security Name"] == source["security_name"]
            assert row["Sector"] == source["sector"]
            assert row["Market Value (INR)"] == source["market_value"]
            assert row["Portfolio Weight"] == source["weight"]
    finally:
        workbook.close()


def test_excel_validation_checks_pass():
    workbook = load_workbook(latest_file("xlsx"), read_only=True, data_only=True)
    try:
        sheet = workbook["Validation Checks"]
        results = {
            sheet.cell(row, 1).value: sheet.cell(row, 2).value
            for row in range(2, sheet.max_row + 1)
        }
        assert results == {
            "Holdings present": "PASS",
            "Portfolio weights": "PASS",
            "Performance record": "PASS",
        }
    finally:
        workbook.close()


def test_powerpoint_has_four_slides_and_fictional_data_disclaimer():
    presentation = Presentation(latest_file("pptx"))
    assert len(presentation.slides) == 4

    text = extract_presentation_text(presentation).lower()
    assert "institutional portfolio report" in text
    assert "illustrative portfolio data" in text
    assert "not actual client or fund reporting" in text
    assert "fictional data" in text


def test_powerpoint_figures_match_database():
    summary, holdings = get_database_data()
    presentation = Presentation(latest_file("pptx"))
    text = extract_presentation_text(presentation)

    expected_values = [
        summary["client"],
        summary["fund"],
        summary["benchmark"],
        f"{summary['total_market_value']:,.0f}",
        str(summary["number_of_holdings"]),
        f"{summary['total_weight']:.2%}",
        f"{summary['fund_return']:.2%}",
        f"{summary['benchmark_return']:.2%}",
        f"{summary['active_return']:.2%}",
    ]
    for value in expected_values:
        assert value in text, f"Missing PowerPoint value: {value}"

    for holding in holdings:
        assert holding["security_id"] in text
        assert holding["security_name"] in text
        assert f"{holding['weight']:.2%}" in text


def test_excel_and_powerpoint_share_run_timestamp():
    assert latest_file("xlsx").stem == latest_file("pptx").stem


def test_successful_pipeline_run_is_audited():
    with sqlite3.connect(DB_PATH) as con:
        row = con.execute("""
            SELECT status, rows_processed, error_message
            FROM pipeline_runs
            WHERE status = 'SUCCEEDED'
            ORDER BY run_id DESC
            LIMIT 1
        """).fetchone()

    assert row is not None
    assert row[0] == "SUCCEEDED"
    assert row[1] == 5
    assert row[2] is None
