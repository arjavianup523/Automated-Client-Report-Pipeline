"""Generate an Excel institutional client report."""

import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "reporting.db"
OUTPUT_DIR = ROOT / "outputs"


def generate_excel_report(
    client_id=1,
    fund_id=1,
    period_id=1,
    run_timestamp=None,
):
    """Generate an Excel report for the selected client, fund, and period."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with sqlite3.connect(DB_PATH) as connection:
        connection.row_factory = sqlite3.Row

        summary = connection.execute(
            """
            SELECT
                c.client_name AS client,
                f.fund_name AS fund,
                f.benchmark_name AS benchmark,
                rp.period_start AS period_start,
                rp.period_end AS period_end,
                COUNT(h.holding_id) AS number_of_holdings,
                SUM(h.market_value) AS total_market_value,
                SUM(h.weight) AS total_weight,
                p.fund_return AS fund_return,
                p.benchmark_return AS benchmark_return,
                p.fund_return - p.benchmark_return AS active_return
            FROM clients c
            JOIN funds f ON f.fund_id = ?
            JOIN reporting_periods rp ON rp.period_id = ?
            JOIN holdings h
                ON h.fund_id = f.fund_id
               AND h.period_id = rp.period_id
            JOIN performance p
                ON p.fund_id = f.fund_id
               AND p.period_id = rp.period_id
            WHERE c.client_id = ?
            GROUP BY
                c.client_name,
                f.fund_name,
                f.benchmark_name,
                rp.period_start,
                rp.period_end,
                p.fund_return,
                p.benchmark_return
            """,
            (fund_id, period_id, client_id),
        ).fetchone()

        holdings = connection.execute(
            """
            SELECT
                security_id,
                security_name,
                sector,
                market_value,
                weight
            FROM holdings
            WHERE fund_id = ? AND period_id = ?
            ORDER BY market_value DESC
            """,
            (fund_id, period_id),
        ).fetchall()

        performance_count = connection.execute(
            """
            SELECT COUNT(*)
            FROM performance
            WHERE fund_id = ? AND period_id = ?
            """,
            (fund_id, period_id),
        ).fetchone()[0]

    if summary is None or not holdings:
        raise ValueError(
            "Required report data is missing for the selected client, "
            "fund, and reporting period."
        )

    if performance_count != 1:
        raise ValueError(
            "Expected exactly one performance record for the selected fund "
            "and reporting period."
        )

    for holding in holdings:
        market_value = holding["market_value"]
        weight = holding["weight"]

        if (
            market_value is None
            or isinstance(market_value, bool)
            or market_value < 0
        ):
            raise ValueError(
                f"Invalid market value for {holding['security_id']}."
            )

        if (
            weight is None
            or not isinstance(weight, (int, float))
            or isinstance(weight, bool)
            or not 0 <= weight <= 1
        ):
            raise ValueError(
                f"Invalid portfolio weight for {holding['security_id']}."
            )

    total_weight = sum(row["weight"] for row in holdings)

    if abs(total_weight - 1.0) > 0.000001:
        raise ValueError(
            f"Report blocked: portfolio weights total "
            f"{total_weight:.2%}, not 100%."
        )

    if summary["fund_return"] is None or summary["benchmark_return"] is None:
        raise ValueError("Fund or benchmark performance data is missing.")

    if run_timestamp is None:
        run_timestamp = datetime.now(timezone.utc).strftime(
            "%Y%m%dT%H%M%SZ"
        )

    output_path = OUTPUT_DIR / f"client_report_{run_timestamp}.xlsx"

    workbook = Workbook()

    # Worksheet 1: Executive summary
    summary_ws = workbook.active
    summary_ws.title = "Executive Summary"
    summary_ws.append(["Metric", "Value"])

    summary_rows = [
        ("Client", summary["client"]),
        ("Fund", summary["fund"]),
        ("Benchmark", summary["benchmark"]),
        ("Period Start", summary["period_start"]),
        ("Period End", summary["period_end"]),
        ("Number of Holdings", summary["number_of_holdings"]),
        ("Total Market Value (INR)", summary["total_market_value"]),
        ("Total Portfolio Weight", total_weight),
        ("Fund Return", summary["fund_return"]),
        ("Benchmark Return", summary["benchmark_return"]),
        ("Active Return", summary["active_return"]),
    ]

    for row in summary_rows:
        summary_ws.append(list(row))

    # Worksheet 2: Holdings
    holdings_ws = workbook.create_sheet("Holdings")
    holdings_ws.append(
        [
            "Security ID",
            "Security Name",
            "Sector",
            "Market Value (INR)",
            "Portfolio Weight",
        ]
    )

    for holding in holdings:
        holdings_ws.append(
            [
                holding["security_id"],
                holding["security_name"],
                holding["sector"],
                holding["market_value"],
                holding["weight"],
            ]
        )

    # Worksheet 3: Validation checks
    checks_ws = workbook.create_sheet("Validation Checks")
    checks_ws.append(["Control", "Result", "Details"])

    checks_ws.append(
        ["Holdings present", "PASS", f"{len(holdings)} holdings loaded"]
    )
    checks_ws.append(
        ["Portfolio weights", "PASS", f"Total weight: {total_weight:.2%}"]
    )
    checks_ws.append(
        ["Performance record", "PASS", "One performance record loaded"]
    )

    # Format worksheet headers and columns.
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(horizontal="center")

        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[letter].width = min(
                max_length + 3, 40
            )

    # Format summary financial values and percentages.
    for row in range(2, summary_ws.max_row + 1):
        metric = summary_ws.cell(row, 1).value
        value_cell = summary_ws.cell(row, 2)

        if metric == "Total Market Value (INR)":
            value_cell.number_format = '"₹"#,##0.00'
        elif metric in (
            "Total Portfolio Weight",
            "Fund Return",
            "Benchmark Return",
            "Active Return",
        ):
            value_cell.number_format = "0.00%"

    # Format holdings financial values and percentages.
    for row in range(2, holdings_ws.max_row + 1):
        holdings_ws.cell(row, 4).number_format = '"₹"#,##0.00'
        holdings_ws.cell(row, 5).number_format = "0.00%"

    workbook.save(output_path)
    workbook.close()

    print(f"Excel report created: {output_path}")
    print("Worksheets: Executive Summary, Holdings, Validation Checks")

    return output_path


if __name__ == "__main__":
    generate_excel_report()
