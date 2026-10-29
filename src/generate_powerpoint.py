
import sqlite3
from pathlib import Path
from datetime import datetime, timezone

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "reporting.db"
OUTPUT_DIR = ROOT / "outputs"

NAVY = RGBColor(31, 78, 121)
BLUE = RGBColor(68, 114, 196)
LIGHT_BLUE = RGBColor(221, 235, 247)
DARK = RGBColor(45, 45, 45)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(0, 120, 80)


def fetch_report_data():
    if not DB_PATH.exists():
        raise FileNotFoundError(f"Reporting database not found: {DB_PATH}")

    with sqlite3.connect(DB_PATH) as connection:
        connection.row_factory = sqlite3.Row

        summary = connection.execute("""
            SELECT
                c.client_name AS client,
                f.fund_name AS fund,
                f.benchmark_name AS benchmark,
                f.base_currency AS currency,
                rp.period_start,
                rp.period_end,
                COUNT(h.holding_id) AS number_of_holdings,
                SUM(h.market_value) AS total_market_value,
                SUM(h.weight) AS total_weight,
                p.fund_return,
                p.benchmark_return,
                p.fund_return - p.benchmark_return AS active_return
            FROM clients c
            JOIN funds f ON f.fund_id = 1
            JOIN reporting_periods rp ON rp.period_id = 1
            JOIN holdings h
                ON h.fund_id = f.fund_id
               AND h.period_id = rp.period_id
            JOIN performance p
                ON p.fund_id = f.fund_id
               AND p.period_id = rp.period_id
            WHERE c.client_id = 1
            GROUP BY
                c.client_name, f.fund_name, f.benchmark_name,
                f.base_currency, rp.period_start, rp.period_end,
                p.fund_return, p.benchmark_return
        """).fetchone()

        holdings = connection.execute("""
            SELECT security_id, security_name, sector,
                   market_value, weight
            FROM holdings
            WHERE fund_id = 1 AND period_id = 1
            ORDER BY market_value DESC
        """).fetchall()

    if summary is None or not holdings:
        raise ValueError("Report generation blocked: required data is missing.")

    total_weight = sum(row["weight"] for row in holdings)

    if abs(total_weight - 1.0) > 0.000001:
        raise ValueError(
            f"Report generation blocked: portfolio weights total "
            f"{total_weight:.2%}, not 100%."
        )

    for holding in holdings:
        if holding["market_value"] is None or holding["market_value"] < 0:
            raise ValueError(
                f"Invalid market value for {holding['security_id']}."
            )

    if summary["fund_return"] is None or summary["benchmark_return"] is None:
        raise ValueError("Report generation blocked: performance data is missing.")

    return summary, holdings


def set_cell_text(cell, text, bold=False, color=DARK, size=12):
    cell.text = str(text)
    cell.vertical_anchor = 1

    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def add_footer(slide, slide_number):
    box = slide.shapes.add_textbox(
        Inches(0.5), Inches(7.12), Inches(12.2), Inches(0.2)
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = (
        "CONFIDENTIALITY: Illustrative sample report — fictional data"
        f"                                                   {slide_number}"
    )
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor(100, 100, 100)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.35), Inches(12), Inches(0.6)
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Aptos Display"
    paragraph.font.size = Pt(27)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY

    if subtitle:
        sub = slide.shapes.add_textbox(
            Inches(0.65), Inches(0.98), Inches(12), Inches(0.35)
        )
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = DARK


def add_metric_card(slide, x, y, width, label, value):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(width), Inches(1.15)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_BLUE
    shape.line.color.rgb = BLUE

    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.vertical_anchor = 3

    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.color.rgb = DARK

    p = tf.add_paragraph()
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = NAVY


def generate_powerpoint():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    summary, holdings = fetch_report_data()

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    # Slide 1: Cover
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = NAVY

    title_box = slide.shapes.add_textbox(
        Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.2)
    )
    p = title_box.text_frame.paragraphs[0]
    p.text = "Institutional Portfolio Report"
    p.font.name = "Aptos Display"
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = WHITE

    details = slide.shapes.add_textbox(
        Inches(0.95), Inches(3.0), Inches(11), Inches(1.4)
    )
    tf = details.text_frame
    tf.text = f"Client: {summary['client']}"
    tf.paragraphs[0].font.size = Pt(19)
    tf.paragraphs[0].font.color.rgb = WHITE

    p = tf.add_paragraph()
    p.text = f"Fund: {summary['fund']}"
    p.font.size = Pt(17)
    p.font.color.rgb = WHITE

    p = tf.add_paragraph()
    p.text = (
        f"Reporting period: {summary['period_start']} to "
        f"{summary['period_end']}"
    )
    p.font.size = Pt(15)
    p.font.color.rgb = WHITE

    disclaimer = slide.shapes.add_textbox(
        Inches(0.95), Inches(6.65), Inches(11.5), Inches(0.3)
    )
    p = disclaimer.text_frame.paragraphs[0]
    p.text = "Illustrative portfolio data — not actual client or fund reporting"
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = WHITE

    # Slide 2: Executive summary
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Executive Summary", "Portfolio overview and performance")

    currency = summary["currency"] or "INR"
    market_value = summary["total_market_value"] or 0

    add_metric_card(
        slide, 0.7, 1.7, 3.7, "PORTFOLIO VALUE",
        f"{currency} {market_value:,.0f}"
    )
    add_metric_card(
        slide, 4.8, 1.7, 3.7, "HOLDINGS",
        str(summary["number_of_holdings"])
    )
    add_metric_card(
        slide, 8.9, 1.7, 3.7, "PORTFOLIO WEIGHT",
        f"{summary['total_weight']:.2%}"
    )

    table = slide.shapes.add_table(
        4, 2, Inches(1.5), Inches(3.5), Inches(10.2), Inches(2.0)
    ).table
    table.columns[0].width = Inches(4.4)
    table.columns[1].width = Inches(5.8)

    rows = [
        ("Reporting benchmark", summary["benchmark"]),
        ("Fund return", f"{summary['fund_return']:.2%}"),
        ("Benchmark return", f"{summary['benchmark_return']:.2%}"),
        ("Active return", f"{summary['active_return']:.2%}"),
    ]

    for r, (label, value) in enumerate(rows):
        set_cell_text(table.cell(r, 0), label, bold=True, size=13)
        set_cell_text(table.cell(r, 1), value, size=13)
        table.cell(r, 0).fill.solid()
        table.cell(r, 0).fill.fore_color.rgb = LIGHT_BLUE

    add_footer(slide, 2)

    # Slide 3: Holdings breakdown
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Portfolio Holdings", "Security-level exposure")

    table = slide.shapes.add_table(
        len(holdings) + 1, 4,
        Inches(0.55), Inches(1.6), Inches(12.2), Inches(4.7)
    ).table

    widths = [1.5, 4.3, 2.8, 1.7]
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)

    headers = ["Security ID", "Security Name", "Sector", "Weight"]
    for col, header in enumerate(headers):
        set_cell_text(table.cell(0, col), header, bold=True, color=WHITE, size=11)
        table.cell(0, col).fill.solid()
        table.cell(0, col).fill.fore_color.rgb = NAVY

    for row_index, holding in enumerate(holdings, start=1):
        values = [
            holding["security_id"],
            holding["security_name"],
            holding["sector"],
            f"{holding['weight']:.2%}",
        ]
        for col, value in enumerate(values):
            set_cell_text(table.cell(row_index, col), value, size=10)
            if row_index % 2 == 0:
                table.cell(row_index, col).fill.solid()
                table.cell(row_index, col).fill.fore_color.rgb = LIGHT_BLUE

    note = slide.shapes.add_textbox(
        Inches(0.65), Inches(6.45), Inches(12), Inches(0.35)
    )
    p = note.text_frame.paragraphs[0]
    p.text = (
        "Weights are checked against a 100% portfolio total before report creation."
    )
    p.font.size = Pt(10)
    p.font.italic = True
    add_footer(slide, 3)

    # Slide 4: Performance interpretation
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Performance Snapshot", "Fund return compared with benchmark")

    fund_return = summary["fund_return"]
    benchmark_return = summary["benchmark_return"]
    active_return = summary["active_return"]

    add_metric_card(
        slide, 0.8, 1.8, 3.6, "FUND RETURN", f"{fund_return:.2%}"
    )
    add_metric_card(
        slide, 4.85, 1.8, 3.6, "BENCHMARK RETURN",
        f"{benchmark_return:.2%}"
    )
    add_metric_card(
        slide, 8.9, 1.8, 3.6, "ACTIVE RETURN", f"{active_return:.2%}"
    )

    interpretation = slide.shapes.add_textbox(
        Inches(1.0), Inches(3.6), Inches(11.2), Inches(1.5)
    )
    tf = interpretation.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    if active_return > 0:
        p.text = (
            f"The fund outperformed its benchmark by {active_return:.2%} "
            "over the selected reporting period."
        )
    elif active_return < 0:
        p.text = (
            f"The fund underperformed its benchmark by {abs(active_return):.2%} "
            "over the selected reporting period."
        )
    else:
        p.text = "The fund return matched the benchmark over the selected period."

    p.font.size = Pt(20)
    p.font.color.rgb = GREEN if active_return >= 0 else RGBColor(180, 50, 50)

    caveat = tf.add_paragraph()
    caveat.text = (
        "Active return is calculated as fund return minus benchmark return. "
        "It is not a measure of risk-adjusted performance."
    )
    caveat.font.size = Pt(12)
    caveat.font.color.rgb = DARK

    add_footer(slide, 4)

    filename = (
        "client_report_"
        + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        + ".pptx"
    )
    output_path = OUTPUT_DIR / filename
    presentation.save(output_path)

    print(f"PowerPoint report created: {output_path}")
    print(f"Slides: {len(presentation.slides)}")
    print("Report data passed the basic portfolio-weight checks.")

    return output_path


if __name__ == "__main__":
    generate_powerpoint()
